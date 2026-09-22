import json
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "vocabulaire.xlsx"
JSON_FILE = BASE_DIR / "vocabulaire.json"

# Si la feuille "Vocabulaire" n'existe pas, le script utilisera
# automatiquement la première feuille du classeur.
SHEET_NAME = "Vocabulaire"


# ============================================================
# FONCTIONS UTILES
# ============================================================

def clean_value(value):
    """
    Nettoie une valeur provenant d'Excel.

    Retourne :
    - une chaîne sans espaces inutiles ;
    - None si la cellule est vide.
    """
    if value is None:
        return None

    value = str(value).strip()

    if not value:
        return None

    return value


def normalize_header(value):
    """
    Normalise un titre de colonne pour faciliter sa détection.
    """
    value = clean_value(value)

    if value is None:
        return ""

    return value.casefold()


def prepare_answers(*values):
    """
    Prépare une liste de réponses :
    - supprime les cellules vides ;
    - supprime les doublons ;
    - conserve l'ordre des traductions.
    """
    answers = []

    for value in values:
        value = clean_value(value)

        if value is None:
            continue

        # Comparaison sans tenir compte des majuscules/minuscules.
        already_exists = any(
            existing.casefold() == value.casefold()
            for existing in answers
        )

        if not already_exists:
            answers.append(value)

    return answers


def is_active(value):
    """
    Vérifie la colonne Actif.

    Sont considérées comme actives :
    - oui
    - o
    - yes
    - y
    - vrai
    - true
    - 1
    - x
    - cellule vide

    Sont considérées comme inactives :
    - non
    - n
    - no
    - faux
    - false
    - 0
    """
    value = clean_value(value)

    # Une cellule Actif vide est considérée comme active.
    if value is None:
        return True

    normalized = value.casefold()

    inactive_values = {
        "non",
        "n",
        "no",
        "faux",
        "false",
        "0",
        "inactif",
        "inactive",
    }

    return normalized not in inactive_values


def find_header_row(sheet, maximum_rows_to_check=20):
    """
    Recherche automatiquement la ligne contenant les titres.

    Le script cherche notamment :
    - Français 1
    - Français 2
    - Français 3
    - Anglais 1
    - Actif
    """
    for row_number in range(
        1,
        min(sheet.max_row, maximum_rows_to_check) + 1,
    ):
        headers = {}

        for column_number in range(1, sheet.max_column + 1):
            cell_value = sheet.cell(
                row=row_number,
                column=column_number,
            ).value

            normalized = normalize_header(cell_value)

            if normalized:
                headers[normalized] = column_number

        if "français 1" in headers and "anglais 1" in headers:
            return row_number, headers

        # Compatibilité si les accents sont absents.
        if "francais 1" in headers and "anglais 1" in headers:
            return row_number, headers

    return None, {}


def get_column_number(headers, *possible_names):
    """
    Retourne le numéro d'une colonne à partir de plusieurs noms possibles.
    """
    for possible_name in possible_names:
        normalized_name = normalize_header(possible_name)

        if normalized_name in headers:
            return headers[normalized_name]

    return None


def get_cell_value(sheet, row_number, column_number):
    """
    Lit une cellule en évitant une erreur si la colonne n'existe pas.
    """
    if column_number is None:
        return None

    return sheet.cell(
        row=row_number,
        column=column_number,
    ).value


# ============================================================
# CONVERSION EXCEL VERS JSON
# ============================================================

def convert_excel_to_json():
    print()
    print("=" * 60)
    print("CONVERSION EXCEL VERS JSON")
    print("=" * 60)

    # --------------------------------------------------------
    # Vérification du fichier Excel
    # --------------------------------------------------------

    if not EXCEL_FILE.exists():
        print()
        print("ERREUR : le fichier Excel est introuvable.")
        print(f"Chemin recherché : {EXCEL_FILE}")
        print()
        print(
            "Place le fichier vocabulaire.xlsx dans le même "
            "dossier que ce script."
        )
        return

    print()
    print(f"Fichier Excel : {EXCEL_FILE}")

    # --------------------------------------------------------
    # Ouverture du classeur
    # --------------------------------------------------------

    try:
        workbook = load_workbook(
            EXCEL_FILE,
            data_only=True,
            read_only=False,
        )

    except PermissionError:
        print()
        print("ERREUR : Python ne peut pas ouvrir le fichier Excel.")
        print("Ferme complètement vocabulaire.xlsx dans Excel.")
        print("Vérifie également la synchronisation OneDrive.")
        return

    except Exception as error:
        print()
        print("ERREUR pendant l'ouverture du fichier Excel :")
        print(error)
        return

    # --------------------------------------------------------
    # Sélection de la feuille
    # --------------------------------------------------------

    print()
    print("Feuilles trouvées dans le classeur :")

    for sheet_name in workbook.sheetnames:
        print(f"  - {sheet_name}")

    if SHEET_NAME in workbook.sheetnames:
        sheet = workbook[SHEET_NAME]
    else:
        sheet = workbook[workbook.sheetnames[0]]

        print()
        print(
            f"ATTENTION : la feuille '{SHEET_NAME}' n'existe pas."
        )
        print(
            f"La feuille '{sheet.title}' sera utilisée automatiquement."
        )

    print()
    print(f"Feuille utilisée : {sheet.title}")
    print(f"Dernière ligne détectée par Excel : {sheet.max_row}")
    print(f"Dernière colonne détectée par Excel : {sheet.max_column}")

    # --------------------------------------------------------
    # Détection automatique des titres
    # --------------------------------------------------------

    header_row, headers = find_header_row(sheet)

    if header_row is None:
        workbook.close()

        print()
        print("ERREUR : impossible de trouver la ligne des titres.")
        print()
        print("Les titres attendus sont :")
        print("  - Français 1")
        print("  - Français 2")
        print("  - Français 3")
        print("  - Anglais 1")
        print("  - Actif")
        return

    print()
    print(f"Ligne des titres détectée : {header_row}")

    french_1_column = get_column_number(
        headers,
        "Français 1",
        "Francais 1",
    )

    french_2_column = get_column_number(
        headers,
        "Français 2",
        "Francais 2",
    )

    french_3_column = get_column_number(
        headers,
        "Français 3",
        "Francais 3",
    )

    english_1_column = get_column_number(
        headers,
        "Anglais 1",
    )

    active_column = get_column_number(
        headers,
        "Actif",
    )

    print()
    print("Colonnes détectées :")
    print(f"  Français 1 : colonne {french_1_column}")
    print(f"  Français 2 : colonne {french_2_column}")
    print(f"  Français 3 : colonne {french_3_column}")
    print(f"  Anglais 1  : colonne {english_1_column}")
    print(f"  Actif      : colonne {active_column}")

    # --------------------------------------------------------
    # Lecture de toutes les lignes
    # --------------------------------------------------------

    words = []

    total_rows_checked = 0
    completely_empty_rows = 0
    inactive_rows = 0
    missing_french_rows = 0
    missing_english_rows = 0

    first_data_row = header_row + 1

    for row_number in range(
        first_data_row,
        sheet.max_row + 1,
    ):
        total_rows_checked += 1

        french_1 = get_cell_value(
            sheet,
            row_number,
            french_1_column,
        )

        french_2 = get_cell_value(
            sheet,
            row_number,
            french_2_column,
        )

        french_3 = get_cell_value(
            sheet,
            row_number,
            french_3_column,
        )

        english_1 = get_cell_value(
            sheet,
            row_number,
            english_1_column,
        )

        active = get_cell_value(
            sheet,
            row_number,
            active_column,
        )

        # Vérification d'une ligne entièrement vide.
        useful_values = [
            french_1,
            french_2,
            french_3,
            english_1,
            active,
        ]

        if all(clean_value(value) is None for value in useful_values):
            completely_empty_rows += 1
            continue

        french_answers = prepare_answers(
            french_1,
            french_2,
            french_3,
        )

        english_answers = prepare_answers(
            english_1,
        )

        # Ligne inactive.
        if not is_active(active):
            inactive_rows += 1
            continue

        # Pas de traduction française.
        if not french_answers:
            missing_french_rows += 1

            print(
                f"Ligne {row_number} ignorée : "
                "aucune traduction française."
            )

            continue

        # Pas de traduction anglaise.
        if not english_answers:
            missing_english_rows += 1

            print(
                f"Ligne {row_number} ignorée : "
                "aucune traduction anglaise."
            )

            continue

        word = {
            "french": french_answers[0],
            "english": english_answers[0],
            "french_answers": french_answers,
            "english_answers": english_answers,
            "category": "Sans catégorie",
        }

        words.append(word)

    workbook.close()

    # --------------------------------------------------------
    # Création du fichier JSON
    # --------------------------------------------------------

    try:
        with JSON_FILE.open(
            "w",
            encoding="utf-8",
        ) as json_file:
            json.dump(
                words,
                json_file,
                ensure_ascii=False,
                indent=4,
            )

    except PermissionError:
        print()
        print("ERREUR : impossible de modifier vocabulaire.json.")
        print("Ferme le fichier JSON s'il est ouvert dans un programme.")
        return

    # --------------------------------------------------------
    # Résumé détaillé
    # --------------------------------------------------------

    print()
    print("=" * 60)
    print("RÉSULTAT")
    print("=" * 60)

    print()
    print(f"Lignes examinées       : {total_rows_checked}")
    print(f"Mots exportés          : {len(words)}")
    print(f"Lignes entièrement vides : {completely_empty_rows}")
    print(f"Lignes inactives       : {inactive_rows}")
    print(f"Français manquant      : {missing_french_rows}")
    print(f"Anglais manquant       : {missing_english_rows}")

    print()
    print(f"Fichier JSON créé : {JSON_FILE}")

    if len(words) == 0:
        print()
        print("ATTENTION : aucun mot n'a été exporté.")

    elif len(words) < 50:
        print()
        print(
            "ATTENTION : peu de mots ont été exportés. "
            "Regarde le bilan ci-dessus pour identifier les lignes ignorées."
        )

    print()
    print("Conversion terminée.")


# ============================================================
# LANCEMENT DU PROGRAMME
# ============================================================

if __name__ == "__main__":
    convert_excel_to_json()